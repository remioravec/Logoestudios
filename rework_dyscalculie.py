# -*- coding: utf-8 -*-
"""Rework SEO de orthophonie/dyscalculie.html (cible data GSC) + DOCX."""
import re, json
FILE="wp-plugin/logopsi-deployer/data/html/orthophonie__dyscalculie.html"
SITE="https://logopsietudes.com"
H1="Bilan dyscalculie : évaluation, test et diagnostic complet"
HERO_SUB=("Un bilan orthophonique complet pour évaluer, tester et diagnostiquer la dyscalculie, ce trouble "
 "durable des apprentissages en mathématiques. Réalisé en ligne par un orthophoniste spécialisé, "
 "premier rendez-vous sous 48h — 150&nbsp;€.")
def L(t,u): return f'<a href="{u}" class="text-primary font-semibold hover:underline">{t}</a>'

LEAD=("<p>La <strong>dyscalculie</strong> est un <strong>trouble</strong> spécifique et durable des "
 "apprentissages qui touche le sens du nombre et le calcul chez l'<strong>enfant</strong> comme chez "
 "l'adulte. Elle est indépendante de l'intelligence. Le <strong>bilan dyscalculie</strong> évalue le "
 "raisonnement <strong>logico-mathématique</strong>, en lien avec le <strong>langage</strong> et le "
 "profil du <strong>patient</strong>, pour poser un diagnostic fiable. Ce suivi s'inscrit dans un "
 + L("écosystème global d'accompagnement", f"{SITE}/") + " (orthophonie, psychologie, soutien scolaire).</p>")

SECTIONS=[
 ("signes","Quels sont les signes de la dyscalculie ?",[
   "<p><strong>Repérer la dyscalculie</strong> chez un enfant, c'est observer des difficultés persistantes en "
   "<strong>mathématiques</strong> malgré les efforts&nbsp;:</p>",
   "<ul class='space-y-2 text-gray-700'>"
   "<li>— un <strong>sens du nombre</strong> fragile (mal évaluer une quantité, comparer deux nombres)&nbsp;;</li>"
   "<li>— des difficultés à mémoriser les <strong>tables</strong> et les faits numériques&nbsp;;</li>"
   "<li>— un <strong>calcul</strong> lent, le recours prolongé au comptage sur les doigts&nbsp;;</li>"
   "<li>— des confusions dans la lecture des nombres et la pose des opérations&nbsp;;</li>"
   "<li>— une difficulté à raisonner sur les problèmes et à se repérer dans le temps et l'espace.</li></ul>",
   "<p>La dyscalculie est souvent associée à d'autres troubles «&nbsp;dys&nbsp;». Si la lecture est aussi "
   "touchée, consultez notre " + L("bilan dyslexie en ligne", f"{SITE}/orthophonie/dyslexie/") + "&nbsp;; "
   "pour l'orthographe, le " + L("bilan dysorthographie en ligne", f"{SITE}/orthophonie/dysorthographie/") + ".</p>",
 ]),
 ("diagnostic","Comment diagnostiquer la dyscalculie ?",[
   "<p>Le diagnostic repose sur un <strong>bilan orthophonique</strong> réalisé par un "
   "<strong>orthophoniste spécialisé</strong> en <strong>mathématiques</strong>. Il évalue les compétences "
   "<strong>logico-mathématiques</strong> (numération, calcul, résolution de problèmes) et les compare aux "
   "attendus de l'âge, à l'aide d'épreuves étalonnées.</p>",
   "<p>Beaucoup de familles cherchent un «&nbsp;<strong>orthophoniste dyscalculie autour de moi</strong>&nbsp;» "
   "et se heurtent à de longs délais. Nos orthophonistes reçoivent <strong>en ligne, partout en France</strong> "
   "— " + L("à Paris", f"{SITE}/orthophonie/villes/paris/") + ", "
   + L("à Marseille", f"{SITE}/orthophonie/villes/dyscalculie-marseille/") + " ou "
   + L("à Nice", f"{SITE}/orthophonie/villes/dyscalculie-nice/") + " — avec un premier bilan sous 48h.</p>",
 ]),
 ("tests","Quels tests pour dépister la dyscalculie ?",[
   "<p>On trouve en ligne de nombreux «&nbsp;<strong>test dyscalculie gratuit</strong>&nbsp;»&nbsp;: ces "
   "questionnaires de <strong>dépistage rapide</strong> repèrent des signes d'alerte, mais restent "
   "<em>indicatifs</em>. Seul un <strong>bilan</strong> avec un professionnel, à l'aide de "
   "<strong>tests étalonnés</strong> (numération, calcul mental, transcodage, comparaison de quantités), "
   "permet de poser un vrai diagnostic.</p>",
 ]),
 ("processus","Comment se déroule le bilan de dyscalculie ?",[
   "<p>Le <strong>bilan complet</strong> se déroule en une à deux séances&nbsp;:</p>",
   "<ol class='space-y-2 text-gray-700 list-decimal list-inside'>"
   "<li><strong>Entretien</strong> avec la famille (parcours scolaire, difficultés observées).</li>"
   "<li>Passation des épreuves logico-mathématiques étalonnées.</li>"
   "<li>Analyse des résultats et <strong>compte rendu</strong> écrit détaillé.</li>"
   "<li>Mise en place d'un plan de rééducation adapté.</li></ol>",
 ]),
 ("accompagner","Comment accompagner un enfant dyscalculique ?",[
   "<p>L'accompagnement associe une <strong>rééducation orthophonique</strong> ciblée sur le nombre et le "
   "calcul, un <strong>soutien scolaire</strong> en mathématiques et des aménagements (calculatrice, "
   "tiers-temps, PAP). L'objectif&nbsp;: reconstruire les bases numériques tout en restaurant la "
   "<strong>confiance</strong> de l'enfant face aux maths.</p>",
   "<p>Un renfort en " + L("soutien scolaire en mathématiques", f"{SITE}/soutien-scolaire/mathematiques/") +
   " consolide les acquis, et notre " + L("prise en charge de l'anxiété scolaire", f"{SITE}/psychologie/anxiete/") +
   " aide lorsque les maths deviennent une source de stress.</p>",
 ]),
]
FAQ=[
 ("À quel âge peut-on diagnostiquer une dyscalculie ?",
  "Un diagnostic fiable est généralement posé à partir du CE2, quand les apprentissages numériques sont "
  "suffisamment avancés pour distinguer une vraie dyscalculie d'un simple retard. Un repérage des signes "
  "d'alerte est toutefois possible bien plus tôt, dès le CP."),
 ("La dyscalculie se soigne-t-elle ?",
  "La dyscalculie ne « disparaît » pas, mais une rééducation orthophonique adaptée permet de très nets "
  "progrès&nbsp;: l'enfant construit des stratégies efficaces et compense durablement ses difficultés."),
 ("Quelle est la différence entre dyscalculie et difficultés passagères en maths ?",
  "Des difficultés passagères se résorbent avec un peu de soutien. La dyscalculie, elle, est durable et "
  "structurelle&nbsp;: les difficultés persistent malgré un accompagnement, sans lien avec le niveau "
  "intellectuel de l'enfant."),
 ("Un orthophoniste peut-il vraiment aider en mathématiques ?",
  "Oui. L'orthophoniste est le professionnel de référence pour les troubles du raisonnement "
  "logico-mathématique&nbsp;: il évalue et rééduque le sens du nombre, le calcul et la résolution de "
  "problèmes, souvent en lien avec le langage."),
 ("Combien coûte un bilan dyscalculie et sous quel délai ?",
  "Chez Logopsi Études, le bilan est à 150&nbsp;€ avec un premier rendez-vous sous 48h, 100&nbsp;% en ligne "
  "et partout en France — une alternative rapide aux longues listes d'attente."),
]

def render_html():
    out=[f'<section class="py-14 px-6 bg-white border-t border-gray-100"><div class="max-w-4xl mx-auto text-gray-600 leading-relaxed text-lg">{LEAD}</div></section>']
    for i,(aid,title,blocks) in enumerate(SECTIONS):
        bg="bg-light" if i%2==0 else "bg-white"
        out.append(f'<section id="{aid}" class="py-16 px-6 {bg} border-t border-gray-100"><div class="max-w-4xl mx-auto">'
                   f'<h2 class="text-3xl font-bold text-gray-900 mb-6">{title}</h2>'
                   f'<div class="space-y-4 text-gray-600 leading-relaxed text-lg">{"".join(blocks)}</div></div></section>')
    faq="".join(f'<div class="bg-white rounded-2xl border border-gray-100 p-6"><h3 class="text-lg font-bold text-gray-900 mb-2">{q}</h3><p class="text-gray-600 leading-relaxed">{a}</p></div>' for q,a in FAQ)
    out.append('<section id="faq" class="py-16 px-6 bg-light border-t border-gray-100"><div class="max-w-4xl mx-auto">'
               '<h2 class="text-3xl font-bold text-gray-900 mb-6">Questions fréquentes sur le bilan de la dyscalculie (FAQ)</h2>'
               f'<div class="space-y-4">{faq}</div></div></section>')
    return "\n".join(out)

def render_jsonld():
    strip=lambda t: re.sub(r'<[^>]+>','',t).replace('&nbsp;',' ').replace('&#39;',"'").replace('&amp;','&')
    d={"@context":"https://schema.org","@type":"FAQPage","mainEntity":[{"@type":"Question","name":strip(q),"acceptedAnswer":{"@type":"Answer","text":strip(a)}} for q,a in FAQ]}
    return '<script type="application/ld+json">'+json.dumps(d,ensure_ascii=False)+'</script>'

def apply_to_page():
    h=open(FILE,encoding="utf-8",errors="replace").read()
    h=re.sub(r'(<h1[^>]*>)[\s\S]*?(</h1>)', lambda m:m.group(1)+H1+m.group(2), h, count=1)
    h=re.sub(r'(<h1[^>]*>[\s\S]*?</h1>\s*<p[^>]*>)[\s\S]*?(</p>)', lambda m:m.group(1)+HERO_SUB+m.group(2), h, count=1)
    start=h.find('<section class="py-20 bg-white">'); end=h.find('<section class="py-20">\n')
    assert start!=-1 and end!=-1 and end>start,(start,end)
    h=h[:start]+render_html()+"\n    "+h[end:]
    if 'FAQPage' not in h:
        hi=h.lower().find('</head>'); h=h[:hi]+render_jsonld()+"\n"+h[hi:]
    h=re.sub(r'<title>.*?</title>','<title>Bilan dyscalculie : évaluation, test et diagnostic | Logopsi Études</title>',h,flags=re.S)
    h=re.sub(r'(<meta name="description" content=")[^"]*(")',r'\g<1>Bilan dyscalculie en ligne : évaluation, test et diagnostic par un orthophoniste spécialisé en mathématiques. Signes, tests, accompagnement. RDV sous 48h — 150€.\2',h,flags=re.I)
    open(FILE,"w",encoding="utf-8").write(h); print("dyscalculie mise à jour")

def make_docx(path):
    from docx import Document
    doc=Document(); clean=lambda t: re.sub(r'<[^>]+>','',t).replace('&nbsp;',' ').replace('&#39;',"'").replace('&amp;','&')
    doc.add_heading(H1,0); doc.add_paragraph(clean(HERO_SUB)); doc.add_paragraph(clean(LEAD))
    for aid,title,blocks in SECTIONS:
        doc.add_heading(clean(title),1)
        for b in blocks:
            if '<ul' in b or '<ol' in b:
                for li in re.findall(r'<li[^>]*>([\s\S]*?)</li>',b): doc.add_paragraph(clean(li).lstrip('— ').strip(),style='List Bullet')
            else: doc.add_paragraph(clean(b))
    doc.add_heading("FAQ",1)
    for q,a in FAQ: doc.add_heading(clean(q),2); doc.add_paragraph(clean(a))
    doc.save(path); print("DOCX:",path)

if __name__=="__main__":
    apply_to_page(); make_docx("/home/user/Logoestudios/Bilan_dyscalculie_Logopsi_Etudes.docx")
